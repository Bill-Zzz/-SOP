#!/usr/bin/env python3
"""生成测试用的小工作簿：一个输入、一个公式、一张图表、一张隐藏表。"""
import sys

import openpyxl
from openpyxl.chart import BarChart, Reference

out = sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/sample.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "假设"
ws["A1"], ws["B1"] = "现金", 6.3
ws["A2"], ws["B2"] = "债务", 35
ws["A3"], ws["B3"] = "净债务", "=B2-B1"
ws["H1"] = "旧出处说明"
for i, v in enumerate([3, 5, 2], start=5):
    ws[f"A{i}"], ws[f"B{i}"] = f"项{i}", v
ch = BarChart()
ch.add_data(Reference(ws, min_col=2, min_row=5, max_row=7))
ws.add_chart(ch, "D5")
h = wb.create_sheet("隐藏计算")
h["A1"] = "=假设!B3*2"
h.sheet_state = "hidden"
wb.save(out)
print("已生成", out)
