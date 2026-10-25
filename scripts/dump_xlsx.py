#!/usr/bin/env python3
"""把 Excel 模型的公式和缓存值一起导出，便于读懂结构。

用法：
  python3 dump_xlsx.py model.xlsx                 # 打印所有工作表概况
  python3 dump_xlsx.py model.xlsx 假设 估值汇总     # 只导出指定工作表
  python3 dump_xlsx.py model.xlsx --csv out_dir   # 每个工作表导出一个 CSV（单元格, 公式/值, 缓存值）
依赖：openpyxl
"""
import csv
import os
import sys

import openpyxl


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    path = sys.argv[1]
    args = sys.argv[2:]
    csv_dir = None
    if "--csv" in args:
        i = args.index("--csv")
        csv_dir = args[i + 1]
        args = args[:i] + args[i + 2:]
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)

    print("== 工作表概况 ==")
    for ws in wf.worksheets:
        print(f"{ws.title}\t{ws.dimensions}\t{ws.sheet_state}\t图表数={len(getattr(ws, '_charts', []))}")
    targets = args or ([] if not csv_dir else [ws.title for ws in wf.worksheets])
    if csv_dir:
        os.makedirs(csv_dir, exist_ok=True)

    for name in targets:
        a, b = wf[name], wv[name]
        rows_out = []
        for row in a.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cached = b[c.coordinate].value
                is_f = isinstance(c.value, str) and c.value.startswith("=")
                rows_out.append((c.coordinate, c.value, cached if is_f else ""))
        if csv_dir:
            fn = os.path.join(csv_dir, f"{name}.csv")
            with open(fn, "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh)
                w.writerow(["cell", "formula_or_value", "cached_value"])
                w.writerows(rows_out)
            print(f"已导出 {fn}（{len(rows_out)} 个单元格）")
        else:
            print(f"\n== {name} ==")
            for coord, v, cached in rows_out:
                s = f"{coord}: {v}"
                if cached != "":
                    s += f"  ->  {cached!r}"
                print(s[:300])


if __name__ == "__main__":
    try:
        main()
    except BrokenPipeError:  # 允许 | head 截断输出
        pass
