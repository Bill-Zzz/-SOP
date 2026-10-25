#!/usr/bin/env python3
"""比较原始与修订后的 xlsx：确认只改了打算改的单元格，图表/隐藏表/公式都还在。

用法：
  python3 xlsx_diff.py original.xlsx revised.xlsx [--values 估值汇总!B14 估值汇总!C14 ...]
输出：
  1) 公式或常量发生变化的单元格（应只出现你打算改的那些）
  2) 工作表隐藏状态变化
  3) 图表/绘图部件数量对比
  4) 指定单元格的重算后数值（原 vs 新）
依赖：openpyxl
"""
import sys
import zipfile

import openpyxl


def parts(path):
    return sorted(n for n in zipfile.ZipFile(path).namelist() if "/charts/" in n or "/drawings/" in n)


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    a_path, b_path = sys.argv[1], sys.argv[2]
    watch = []
    if "--values" in sys.argv:
        watch = sys.argv[sys.argv.index("--values") + 1:]
    a = openpyxl.load_workbook(a_path)
    b = openpyxl.load_workbook(b_path)
    n = 0
    print("== 内容变化（公式/常量）==")
    for wa in a.worksheets:
        if wa.title not in b.sheetnames:
            print(f"缺少工作表：{wa.title}")
            continue
        wb_ = b[wa.title]
        if wa.sheet_state != wb_.sheet_state:
            print(f"[隐藏状态] {wa.title}: {wa.sheet_state} -> {wb_.sheet_state}")
        for row in wa.iter_rows():
            for c in row:
                v2 = wb_[c.coordinate].value
                if c.value != v2:
                    n += 1
                    print(f"{wa.title}!{c.coordinate}: {str(c.value)[:60]!r} -> {str(v2)[:60]!r}")
    print(f"共 {n} 处变化")
    pa, pb = parts(a_path), parts(b_path)
    print(f"\n== 图表/绘图部件 == 原 {len(pa)} 个，新 {len(pb)} 个" + ("" if pa == pb else "  ⚠ 不一致"))
    if watch:
        av = openpyxl.load_workbook(a_path, data_only=True)
        bv = openpyxl.load_workbook(b_path, data_only=True)
        print("\n== 关键数值（重算后）==")
        for ref in watch:
            s, c = ref.split("!")
            print(f"{ref}: {av[s][c].value} -> {bv[s][c].value}")


if __name__ == "__main__":
    main()
